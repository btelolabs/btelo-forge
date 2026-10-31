"""Checks xlsx files written by internal/office/xlsx with openpyxl, an independent reader.

Run with the office-check venv from the repository's top folder:

    python tools/officecheck/sheet_check.py

It makes a rich workbook with openpyxl, has the Go code open, edit and save it (and write blank and made files of its
own), then confirms that the edits took and nothing else changed.
"""
import datetime
import os
import subprocess
import sys
import tempfile
import zipfile

import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

failures = []


def check(cond, what):
    print(("ok   " if cond else "FAIL ") + what)
    if not cond:
        failures.append(what)


def make_rich(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Name", "Qty", "Price", "Total", "When"])
    for i, (n, q, p) in enumerate([("Alpha", 3, 1.5), ("Beta", 5, 2.25), ("Gamma", 7, 9.75), ("Delta", 2, 4)], start=2):
        ws.append([n, q, p, f"=B{i}*C{i}", datetime.date(2024, 1, i)])
    ws["F1"] = "Sum"
    ws["F2"] = "=SUM(D2:D5)"
    ws["A1"].font = Font(name="Arial", size=14, bold=True, color="FF0000")
    ws["B1"].fill = PatternFill("solid", fgColor="FFFF00")
    ws["C2"].number_format = '#,##0.00;[Red]-#,##0.00'
    ws["D2"].border = Border(left=Side(style="thin"), bottom=Side(style="double", color="0000FF"))
    ws["A2"].alignment = Alignment(horizontal="center", wrap_text=True)
    ws["E2"].number_format = "yyyy-mm-dd"
    ws.column_dimensions["A"].width = 22
    ws.row_dimensions[3].height = 33
    ws.merge_cells("H1:I2")
    ws["H1"] = "Merged"
    ws.freeze_panes = "B2"
    ws.sheet_properties.tabColor = "00B050"
    ws.conditional_formatting.add("B2:B5", CellIsRule(operator="greaterThan", formula=["4"], fill=PatternFill(bgColor="FFC7CE")))
    dv = DataValidation(type="list", formula1='"x,y,z"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("G2:G5")
    ws["A5"].comment = Comment("A note", "Tester")
    ws["A4"].hyperlink = "https://example.com/"
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=5), titles_from_data=True)
    ws.add_chart(chart, "K2")
    t = wb.create_sheet("Table")
    t.append(["Key", "Value"])
    t.append(["a", 1])
    t.append(["b", 2])
    tab = Table(displayName="Things", ref="A1:B3")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    t.add_table(tab)
    t["D1"] = "=Data!F2*2"
    wb.defined_names["Rate"] = DefinedName("Rate", attr_text="Data!$C$2")
    h = wb.create_sheet("Hidden stuff")
    h["A1"] = "=Rate+1"
    h.sheet_state = "hidden"
    wb.save(path)


def parts(path):
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist()}


def main():
    root = os.getcwd()
    with tempfile.TemporaryDirectory() as d:
        make_rich(os.path.join(d, "in.xlsx"))
        env = dict(os.environ, FORGE_XLSX_CHECK=d)
        out = subprocess.run(["go", "test", "./internal/office/xlsx", "-run", "TestOfficeCheck", "-count=1"], cwd=root, env=env, capture_output=True, text=True)
        print(out.stdout[-2000:], out.stderr[-2000:])
        check(out.returncode == 0, "the Go side ran")

        wb = openpyxl.load_workbook(os.path.join(d, "blank.xlsx"))
        check(wb.sheetnames == ["Sheet1"], "blank: one sheet")

        wb = openpyxl.load_workbook(os.path.join(d, "made.xlsx"))
        vals = openpyxl.load_workbook(os.path.join(d, "made.xlsx"), data_only=True)
        ws, vs = wb["Sheet1"], vals["Sheet1"]
        check(ws["A1"].value == "Hello 世界", "made: text")
        check(ws["B1"].value == 2.5, "made: number")
        check(ws["C1"].value == "=B1*2", "made: formula")
        check(vs["C1"].value == 5, "made: cached value")
        check(ws["A1"].font.b and ws["A1"].font.sz == 14 and ws["A1"].font.color.rgb == "FF112233", "made: font")
        check(ws["A1"].fill.fgColor.rgb == "FFFFFF00", "made: fill")
        check(ws["B1"].number_format == "0.00%", "made: number format")
        check(ws["D2"].border.left.style == "thin", "made: border")
        check(ws["D2"].alignment.horizontal == "center" and ws["D2"].alignment.wrap_text, "made: alignment")
        check("A3:B4" in [str(m) for m in ws.merged_cells.ranges], "made: merge")
        check(ws.freeze_panes == "B2", "made: freeze")
        check(ws.sheet_properties.tabColor.rgb == "FF00B050", "made: tab colour")
        check(abs(ws.column_dimensions["F"].width - 30) < 0.01, "made: width")
        check(ws.row_dimensions[6].height == 40, "made: height")
        check(ws.row_dimensions[7].hidden, "made: hidden row")
        check(wb.sheetnames == ["Sheet1", "Second"], "made: sheets")
        check(vals["Second"]["A1"].value == 10, "made: sheet reference value")
        check(ws.auto_filter.ref == "H1:I4", "made: filter")

        before = openpyxl.load_workbook(os.path.join(d, "in.xlsx"))
        after = openpyxl.load_workbook(os.path.join(d, "out.xlsx"))
        vals = openpyxl.load_workbook(os.path.join(d, "out.xlsx"), data_only=True)
        a, b = before["Data"], after["Data"]
        check(after.sheetnames == before.sheetnames, "out: sheets")
        check(b["B3"].value == 50, "out: the edit took")
        check(vals["Data"]["D3"].value == 112.5, "out: dependents recalculated")
        check(vals["Table"]["D1"].value == vals["Data"]["F2"].value * 2, "out: across sheets")
        check(b["A1"].font.b and b["A1"].font.name == "Arial" and b["A1"].font.color.rgb == "00FF0000", "out: font unchanged")
        check(b["C2"].number_format == a["C2"].number_format, "out: number format unchanged")
        check(b["D2"].border.bottom.style == "double", "out: border unchanged")
        check(b.freeze_panes == "B2", "out: freeze unchanged")
        check("H1:I2" in [str(m) for m in b.merged_cells.ranges], "out: merge unchanged")
        check(len(b.conditional_formatting) == 1 and len(b.data_validations.dataValidation) == 1, "out: rules kept")
        check(b["A5"].comment is not None, "out: comment kept")
        check(b["A4"].hyperlink is not None, "out: link kept")
        check(len(b._charts) == 1, "out: chart kept")
        check("Things" in after["Table"].tables, "out: table kept")
        check("Rate" in after.defined_names, "out: name kept")
        check(after["Hidden stuff"].sheet_state == "hidden", "out: hidden sheet kept")
        p, q = parts(os.path.join(d, "in.xlsx")), parts(os.path.join(d, "out.xlsx"))
        changed = sorted(n for n in p if n in q and p[n] != q[n])
        gone = sorted(set(p) - set(q))
        print("changed parts:", changed, "gone:", gone)
        check(set(changed) <= {"[Content_Types].xml", "xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml", "xl/worksheets/sheet3.xml", "xl/workbook.xml", "xl/sharedStrings.xml"}, "out: other parts identical")
        check(not gone, "out: no part lost")

        # Rows inserted: everything follows.
        wb = openpyxl.load_workbook(os.path.join(d, "rows.xlsx"))
        vals = openpyxl.load_workbook(os.path.join(d, "rows.xlsx"), data_only=True)
        ws = wb["Data"]
        check(ws["D6"].value == "=B6*C6", "rows: formula moved")
        check(ws["F2"].value == "=SUM(D2:D7)", "rows: range grew")
        check(wb["Table"]["D1"].value == "=Data!F2*2", "rows: other sheet follows")
        check("H1:I2" in [str(m) for m in ws.merged_cells.ranges], "rows: merge above stays")
        check(str(next(iter(ws.conditional_formatting)).sqref) == "B2:B7", "rows: rule's range grew")
        check(ws.row_dimensions[5].height == 33, "rows: height moved")
        check(ws["A7"].comment is not None, "rows: comment moved")
        check(ws._charts[0].anchor._from.row == 1, "rows: chart above stays")
        check(vals["Data"]["F2"].value == 92, "rows: value")
    if failures:
        print(len(failures), "failed")
        sys.exit(1)
    print("all good")


main()
