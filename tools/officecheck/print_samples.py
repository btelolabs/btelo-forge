"""Makes rich Office files for checking printing: mixed page sizes, tables, shapes, gradients, pictures, print setups.

Usage: python print_samples.py OUT_DIR
Needs python-docx, openpyxl, python-pptx and Pillow (the office-check venv).
"""
import os
import sys

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from pptx import Presentation
from pptx.dml.color import RGBColor as PptColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.util import Emu, Inches as PIn, Pt as PPt


def picture(path, jpeg=False):
    img = Image.new("RGB", (800, 500), (40, 90, 170))
    d = ImageDraw.Draw(img)
    for i in range(0, 800, 40):
        d.rectangle([i, 0, i + 20, 500], fill=(70 + i // 8, 140, 200 - i // 10))
    d.ellipse([250, 100, 550, 400], fill=(250, 210, 60))
    img.save(path, quality=90)


def word(out, pic):
    d = Document()
    d.add_heading("打印测试 Print test", 0)
    p = d.add_paragraph("A paragraph of ordinary words, with ")
    p.add_run("bold").bold = True
    p.add_run(", ")
    p.add_run("italic").italic = True
    r = p.add_run(" and coloured words")
    r.font.color.rgb = RGBColor(0xC0, 0x20, 0x20)
    p.add_run(". 中文段落也要是矢量文字，可以搜索和复制。")
    t = d.add_table(rows=3, cols=3)
    t.style = "Table Grid"
    for i in range(3):
        for j in range(3):
            t.cell(i, j).text = f"R{i + 1} C{j + 1}"
    d.add_paragraph()
    d.add_paragraph("Bullet one", style="List Bullet")
    d.add_paragraph("Bullet two", style="List Bullet")
    d.add_picture(pic, width=Inches(3))
    s = d.add_section(WD_SECTION.NEW_PAGE)
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = s.page_height, s.page_width
    d.add_heading("Landscape section 横向小节", 1)
    d.add_paragraph("This page is wider than it is high. " * 20)
    s = d.add_section(WD_SECTION.NEW_PAGE)
    s.orientation = WD_ORIENT.PORTRAIT
    s.page_width, s.page_height = Inches(5.83), Inches(8.27)
    d.add_heading("A5 section", 1)
    d.add_paragraph("A smaller sheet.")
    d.sections[0].header.paragraphs[0].text = "The header"
    d.save(os.path.join(out, "word.docx"))


def sheet(out):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="2F5597")
    thin = Side(style="thin", color="000000")
    ws.append(["Region", "地区", "Q1", "Q2", "Q3", "Q4", "Total"])
    for c in ws[1]:
        c.font, c.fill = head, fill
        c.alignment = Alignment(horizontal="center")
    for i in range(2, 120):
        ws.append([f"Region {i}", "华东", i * 3, i * 5, i * 7, i * 2, None])
        ws.cell(i, 7).value = f"=SUM(C{i}:F{i})"
        for c in ws[i]:
            c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.column_dimensions["A"].width = 18
    ws.print_title_rows = "1:1"
    ws.oddHeader.center.text = "&A"
    ws.oddFooter.right.text = "Page &P of &N"
    ws.print_options.gridLines = False
    ws2 = wb.create_sheet("Wide")
    for r in range(1, 12):
        for c in range(1, 30):
            ws2.cell(r, c).value = r * c
    ws2.page_setup.orientation = "landscape"
    ws2.page_setup.paperSize = ws2.PAPERSIZE_A4
    ws2.sheet_properties.pageSetUpPr.fitToPage = True
    ws2.page_setup.fitToHeight = 0
    ws2.print_options.gridLines = True
    ws2.print_options.headings = True
    ws3 = wb.create_sheet("Area")
    for r in range(1, 40):
        ws3.append([f"line {r}", r, r * r])
    ws3.print_area = "A5:C20"
    wb.save(os.path.join(out, "sheet.xlsx"))


def slides(out, pic, jpg):
    p = Presentation()
    p.slide_width, p.slide_height = PIn(13.333), PIn(7.5)
    s = p.slides.add_slide(p.slide_layouts[0])
    s.shapes.title.text = "打印测试 Print test"
    s.placeholders[1].text = "Vector slides"
    s.notes_slide.notes_text_frame.text = "Speaker notes for the first slide. 第一页的备注。"
    s = p.slides.add_slide(p.slide_layouts[6])
    a = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PIn(0.5), PIn(0.5), PIn(4), PIn(2))
    a.fill.gradient()
    a.fill.gradient_angle = 30
    a.fill.gradient_stops[0].color.rgb = PptColor(0xE0, 0x30, 0x30)
    a.fill.gradient_stops[1].color.rgb = PptColor(0x30, 0x30, 0xE0)
    a.text_frame.text = "Gradient 渐变"
    b = s.shapes.add_shape(MSO_SHAPE.OVAL, PIn(5), PIn(0.5), PIn(3), PIn(2))
    b.fill.solid()
    b.fill.fore_color.rgb = PptColor(0x30, 0xB0, 0x60)
    b.rotation = 25
    b.text_frame.text = "Turned 旋转"
    c = s.shapes.add_shape(MSO_SHAPE.STAR_5_POINT, PIn(9), PIn(0.5), PIn(2.5), PIn(2.5))
    c.fill.solid()
    c.fill.fore_color.rgb = PptColor(0xF0, 0xC0, 0x20)
    c.line.color.rgb = PptColor(0x80, 0x40, 0x00)
    c.line.width = PPt(3)
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, PIn(0.5), PIn(3), PIn(5), PIn(3.5))
    ln.line.color.rgb = PptColor(0, 0, 0)
    ln.line.width = PPt(2)
    ln.line.dash_style = 4
    pc = s.shapes.add_picture(jpg, PIn(0.5), PIn(4), width=PIn(3))
    pc.crop_left, pc.crop_right = 0.2, 0.2
    t = s.shapes.add_table(3, 3, PIn(5), PIn(4), PIn(5), PIn(2)).table
    for i in range(3):
        for j in range(3):
            t.cell(i, j).text = f"{i},{j}"
    s = p.slides.add_slide(p.slide_layouts[5])
    s.shapes.title.text = "Hidden slide"
    s._element.set("show", "0")
    p.save(os.path.join(out, "slides.pptx"))


def main():
    out = sys.argv[1]
    os.makedirs(out, exist_ok=True)
    pic, jpg = os.path.join(out, "pic.png"), os.path.join(out, "pic.jpg")
    picture(pic)
    picture(jpg)
    word(out, pic)
    sheet(out)
    slides(out, pic, jpg)


if __name__ == "__main__":
    main()
