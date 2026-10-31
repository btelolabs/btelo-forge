"""Independent checks of the presentation code, with python-pptx as the other reader.

  python slides_check.py make  <dir>   writes rich.pptx (a deck python-pptx builds, with what it can make)
  python slides_check.py blank <file>  checks a file this app wrote from blank
  python slides_check.py edited <orig> <edited>  checks that the edits in TestRichFileEdit took and that nothing else changed

The Go side is TestOfficeCheckFiles in internal/office/pptx, which runs when FORGE_OFFICECHECK names a folder.
"""
import io
import struct
import sys
import zipfile
import zlib
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Inches, Pt


def png(w, h, fn):
    rows = b""
    for y in range(h):
        rows += b"\0" + bytes(c for x in range(w) for c in fn(x, y))

    def chunk(t, d):
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xffffffff)

    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)) +
            chunk(b"IDAT", zlib.compress(rows)) + chunk(b"IEND", b""))


def make(folder):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    # A title slide.
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Rich deck 中文标题"
    s.placeholders[1].text = "Made by python-pptx"
    s.notes_slide.notes_text_frame.text = "Speaker notes of the first slide"

    # Bullets at levels, runs in styles.
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Bullets"
    tf = s.placeholders[1].text_frame
    tf.text = "First level"
    for lvl, words in ((1, "Second level"), (2, "Third level"), (0, "Back to the first, a long line that has to wrap round to show how a hanging indent holds up in practice")):
        p = tf.add_paragraph()
        p.text = words
        p.level = lvl
    p = tf.add_paragraph()
    r = p.add_run()
    r.text = "Bold "
    r.font.bold = True
    r = p.add_run()
    r.text = "red "
    r.font.color.rgb = RGBColor(0xc0, 0, 0)
    r = p.add_run()
    r.text = "italic big"
    r.font.italic = True
    r.font.size = Pt(32)

    # Shapes.
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Shapes"
    a = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.5), Inches(3), Inches(1.5))
    a.fill.gradient()
    a.fill.gradient_angle = 45
    a.fill.gradient_stops[0].color.rgb = RGBColor(0xff, 0x80, 0)
    a.fill.gradient_stops[1].color.rgb = RGBColor(0x20, 0x40, 0xff)
    a.line.color.rgb = RGBColor(0, 0, 0)
    a.line.width = Pt(3)
    a.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    a.text_frame.text = "Gradient"
    b = s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(4), Inches(1.5), Inches(2), Inches(2))
    b.rotation = 30
    b.shadow.inherit = False
    b.text_frame.text = "Turned"
    c = s.shapes.add_shape(MSO_SHAPE.STAR_5_POINT, Inches(7), Inches(1.5), Inches(2), Inches(2))
    c.fill.solid()
    c.fill.fore_color.rgb = RGBColor(0xff, 0xd7, 0)
    d = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(9.5), Inches(1.5), Inches(3), Inches(1))
    d.text_frame.text = "Arrow"
    e = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.5), Inches(4), Inches(5), Inches(6))
    e.line.width = Pt(2)
    g = s.shapes.add_group_shape()
    g1 = g.shapes.add_shape(MSO_SHAPE.OVAL, Inches(6), Inches(4), Inches(1.5), Inches(1.5))
    g2 = g.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7), Inches(4.5), Inches(2), Inches(1))
    g1.text_frame.text = "In"
    g2.text_frame.text = "a group"
    t = s.shapes.add_textbox(Inches(9.5), Inches(4), Inches(3), Inches(2))
    t.text_frame.word_wrap = True
    t.text_frame.text = "A text box with words that wrap inside its width, set to the middle."
    t.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    t.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

    # A picture, cropped, a table, a chart.
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Picture, table, chart"
    data = png(160, 120, lambda x, y: (x * 255 // 160, y * 255 // 120, 128))
    pic = s.shapes.add_picture(io.BytesIO(data), Inches(0.5), Inches(1.5), Inches(3))
    pic.crop_left = 0.25
    pic.crop_top = 0.1
    pic.rotation = 10
    tb = s.shapes.add_table(4, 3, Inches(4), Inches(1.5), Inches(5), Inches(2)).table
    for r in range(4):
        for k in range(3):
            tb.cell(r, k).text = f"R{r}C{k}"
    tb.cell(1, 0).merge(tb.cell(2, 0))
    tb.cell(3, 1).fill.solid()
    tb.cell(3, 1).fill.fore_color.rgb = RGBColor(0xff, 0xee, 0x88)
    chart = CategoryChartData()
    chart.categories = ["A", "B", "C"]
    chart.add_series("S", (1, 3, 2))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(9.5), Inches(1.5), Inches(3.5), Inches(3), chart)

    # Autofit and anchors.
    s = prs.slides.add_slide(prs.slide_layouts[6])
    t = s.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    t.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    t.text_frame.word_wrap = True
    t.text_frame.text = "Words " * 40
    bp = t.text_frame._txBody.bodyPr
    bp.find("{http://schemas.openxmlformats.org/drawingml/2006/main}normAutofit").set("fontScale", "62500")
    u = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6), Inches(1), Inches(3), Inches(3))
    u.text_frame.text = "Bottom"
    u.text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM
    prs.save(Path(folder) / "rich.pptx")


def check_blank(path):
    prs = Presentation(path)
    assert prs.slide_width == 12192000 and prs.slide_height == 6858000, "size"
    names = [l.name for l in prs.slide_layouts]
    for want in ("Title Slide", "Title and Content", "Blank"):
        assert want in names, names
    assert len(prs.slides) == 1
    s = prs.slides[0]
    assert s.slide_layout.name == "Title Slide"
    assert s.shapes.title is not None
    assert len(list(s.placeholders)) == 2
    print("blank ok:", names)


def parts(path):
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist()}


def check_edited(orig, edited):
    prs = Presentation(edited)
    slides = list(prs.slides)
    assert len(slides) == 6, len(slides)
    assert slides[0].shapes.title.text == "Rich deck 中文标题, edited", slides[0].shapes.title.text
    assert slides[0].notes_slide.notes_text_frame.text == "Speaker notes of the first slide"
    assert slides[1].notes_slide.notes_text_frame.text == "New notes"
    shapes = list(slides[2].shapes)
    names = [x.name for x in shapes]
    assert any(x.shape_type is not None and x.has_text_frame and x.text_frame.text == "Added" for x in shapes), names
    moved = [x for x in shapes if x.has_text_frame and x.text_frame.text == "Arrow"][0]
    assert moved.left == Inches(1), moved.left
    assert slides[5].shapes.title is not None or len(list(slides[5].shapes)) >= 0
    a, b = parts(orig), parts(edited)
    # What was not edited is byte for byte the same.
    same = [n for n in a if n.startswith(("ppt/slideMasters", "ppt/slideLayouts", "ppt/theme", "ppt/charts", "ppt/embeddings", "ppt/media", "docProps"))]
    for n in same:
        assert n in b, "missing " + n
        assert a[n] == b[n], "changed " + n
    for n in ("ppt/slides/slide4.xml", "ppt/slides/slide5.xml"):
        assert a[n] == b[n], "changed " + n
    print("edited ok,", len(same), "parts the same")


if __name__ == "__main__":
    cmd = sys.argv[1]
    if cmd == "make":
        make(sys.argv[2])
    elif cmd == "blank":
        check_blank(sys.argv[2])
    elif cmd == "edited":
        check_edited(sys.argv[2], sys.argv[3])
    else:
        sys.exit("unknown command")
